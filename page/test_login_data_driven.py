import os
import pytest
from openpyxl import load_workbook


class Test_DataDriven:
    """数据驱动测试类"""

    @staticmethod
    def load_test_data(file_name="login_data.xlsx", sheet_name="Sheet1"):
        """
        从Excel文件加载测试数据
        :param file_name: Excel文件名
        :param sheet_name: 工作表名称
        :return: 测试数据列表
        """
        base_dir = os.path.dirname(os.path.abspath(__file__))
        file_path = os.path.join(base_dir, "..", "test_data", file_name)
        test_data = []

        try:
            workbook = load_workbook(file_path)
            sheet = workbook[sheet_name] if sheet_name in workbook.sheetnames else workbook.active

            for row in sheet.iter_rows(min_row=2, values_only=True):
                if not row[12]:  # phone 字段是第13列（索引12）
                    continue

                test_data.append({
                    "code": row[0],
                    "dynamicPassword": row[1],
                    "graphicsKey": row[2],
                    "imgCode": row[3],
                    "iosOrAndroid": row[4] if row[4] is not None else 0,
                    "isApp": row[5] if row[5] is not None else 0,
                    "lastCity": row[6],
                    "lastProvince": row[7],
                    "openType": row[8],
                    "parentId": row[9],
                    "parentLink": row[10],
                    "parentType": row[11] if row[11] is not None else 0,
                    "phone": row[12],
                    "password": row[13],
                    "rid": row[14],
                    "temporaryToken": row[15],
                    "ydToken": row[16]
                })

            workbook.close()
            return test_data

        except FileNotFoundError:
            print(f"文件 '{file_path}' 不存在")
        except Exception as e:
            print(f"读取文件 '{file_path}' 时发生错误: {e}")

        return test_data

    def test_load_login_data(self):
        """测试加载登录数据"""
        data = self.load_test_data()
        assert data is not None, "数据加载失败"
        assert isinstance(data, list), "数据格式错误"
        
        print(f"\n{'='*60}")
        print(f"成功加载 {len(data)} 条测试数据")
        print(f"{'='*60}")
        
        for i, test_case in enumerate(data, 1):
            print(f"\n【测试用例 {i}】")
            print(f"  phone:    {test_case.get('phone', '')}")
            print(f"  password: {test_case.get('password', '')}")
            print(f"  code:     {test_case.get('code', '')}")
            print(f"  graphicsKey: {test_case.get('graphicsKey', '')}")
            print(f"  imgCode:      {test_case.get('imgCode', '')}")
            print(f"  iosOrAndroid: {test_case.get('iosOrAndroid', '')}")
            print(f"  isApp:        {test_case.get('isApp', '')}")
            print(f"  lastCity:    {test_case.get('lastCity', '')}")
            print(f"  lastProvince: {test_case.get('lastProvince', '')}")
            print(f"  openType:    {test_case.get('openType', '')}")
            print(f"  parentId:    {test_case.get('parentId', '')}")
            print(f"  parentLink:  {test_case.get('parentLink', '')}")
            print(f"  parentType:  {test_case.get('parentType', '')}")
            print(f"  rid:         {test_case.get('rid', '')}")
            print(f"  temporaryToken: {test_case.get('temporaryToken', '')}")
            print(f"  ydToken:        {test_case.get('ydToken', '')}")
            print(f"  dynamicPassword: {test_case.get('dynamicPassword', '')}")
        print(f"\n{'='*60}")

    def test_data_file_exists(self):
        """测试数据文件是否存在"""
        base_dir = os.path.dirname(os.path.abspath(__file__))
        file_path = os.path.join(base_dir, "..", "test_data", "login_data.xlsx")
        assert os.path.exists(file_path), f"数据文件不存在: {file_path}"